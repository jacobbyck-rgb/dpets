"""Pull reference data out of Diamond's ticket workbook (Google Sheets export)
and write it as JSON for the playbook page.

Usage: python3 tools/extract_data.py "Diamond Pet foods - Revised.xlsx" > data.json
"""
import json
import sys
from collections import Counter
from datetime import datetime

import openpyxl


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def table(ws):
    rows = list(ws.iter_rows(values_only=True))
    hdr = [clean(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(clean(v) for v in r):
            continue
        out.append({h: clean(v) for h, v in zip(hdr, r) if h})
    return out


def main(path):
    wb = openpyxl.load_workbook(path)

    depots = [
        {"code": r["Code"], "facility": r["Facility"], "city": r["City Go to WHS"], "contact": r.get("Contact Info", "")}
        for r in table(wb["Costco Locations PO reference"]) if r["Code"]
    ]
    adww_map = [
        {"region": r["Region"], "facility": r["Facility"], "code": r["Code"], "contact": r["Contact_ID"], "notes": r["Notes"]}
        for r in table(wb["ADWW Code Map"])
    ]
    contacts = {r["Contact_ID"]: {"name": r["Contact Name"], "email": r["Emails"]} for r in table(wb["ADWW Contacts"])}
    warehouses = [
        {
            "city": r["City1"], "state": r["State"], "zip": r["Zip"], "distance": r["Distance"],
            "code": r["Warehouse_Code/ID"], "name": r["Warehouse Name"], "hours": r["Hours"],
            "address": r["Address"], "contact": r["Contact"], "delivery": r["$"],
            "rates": r["Rate_per_Pallet"], "storage": r["Crossdock_Storage_Rates"], "notes": r["Internal_Notes"],
        }
        for r in table(wb["Go to Warehouses"]) if r["City1"]
    ]
    skus = [
        {"costco": r["Costco #"], "diamond": r["Diamond#"], "size": r["Sz"], "name": r["Item Description"],
         "pltQty": r["Plt Qty"], "pltWeight": r["Plt Weight"]}
        for r in table(wb["Diamond Costco POs"]) if r["Item Description"]
    ]
    carriers = sorted({clean(c[0].value) for c in wb["Aux"].iter_rows(min_row=2) if clean(c[0].value)})
    origins = [f'{r["City"]} {r["State"]} {r["Zip"]}' for r in table(wb["Diamond Locations"])]

    tickets = [t for t in table(wb["Ticket System"]) if t["WarehouseNow_PO"]]
    inv = [float(t["Invoice_Total"]) for t in tickets
           if t["Invoice_Total"].replace(".", "", 1).isdigit() and float(t["Invoice_Total"]) > 0]
    stats = {
        "tickets": len(tickets),
        "firstDate": min(t["Date"] for t in tickets if t["Date"]),
        "lastDate": max(t["Date"] for t in tickets if t["Date"]),
        "reason": Counter(t["Reason_Refused"] or "(blank)" for t in tickets).most_common(),
        "disposition": Counter(t["Disposition"] or "(blank)" for t in tickets).most_common(),
        "status": Counter(t["Delivery_Status"] or "(blank)" for t in tickets).most_common(),
        "destination": Counter(t["Destination"] for t in tickets if t["Destination"]).most_common(10),
        "origin": Counter(t["Origin"] for t in tickets if t["Origin"]).most_common(),
        "carrier": Counter(t["Carrier"] for t in tickets if t["Carrier"]).most_common(8),
        "invoiced": len(inv),
        "invoiceSum": round(sum(inv), 2),
        "invoiceAvg": round(sum(inv) / len(inv), 2) if inv else 0,
    }
    json.dump({"depots": depots, "adwwMap": adww_map, "contacts": contacts, "warehouses": warehouses,
               "skus": skus, "carriers": carriers, "origins": origins, "stats": stats}, sys.stdout, indent=1)


if __name__ == "__main__":
    main(sys.argv[1])
